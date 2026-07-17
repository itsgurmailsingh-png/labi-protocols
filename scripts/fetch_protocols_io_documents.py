"""
fetch_protocols_io_documents.py

Recovers content for protocols_io protocols currently showing zero steps.
Found via the "why is this empty" investigation: 500 of 776 zero-step
protocols_io records have a real attached document in the cached API
response (data/raw/protocols_io/*.json -> protocol.documents[]) that was
never fetched or examined — the protocol author put the actual procedure
in a PDF/DOCX attachment instead of protocols.io's structured step editor.
Confirmed the download URL works directly:
https://content.protocols.io/files/{key}.{ext} -> 200, real file.

For PDF/DOCX (the vast majority — 548 of ~570 total attachments): download,
extract text, and if it looks like real step content, populate steps_raw
so these protocols stop being blank. Every attachment (any file type) also
gets recorded as a downloadable "media" entry regardless of whether text
extraction succeeds — even if we can't parse steps out of it, a real
protocol shouldn't show as empty when there's a document a scientist could
open and follow.

Usage:
    python3 scripts/fetch_protocols_io_documents.py --dry-run
    python3 scripts/fetch_protocols_io_documents.py
"""

import argparse
import json
import pathlib
import re
import time

import requests

try:
    import docx as docx_lib
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import pypdf
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

CACHE_DIR = pathlib.Path("data/raw/protocols_io")
NORM_DIR = pathlib.Path("data/sources/protocols_io/normalised")
RAW_DIR = pathlib.Path("data/sources/protocols_io/raw")
MEDIA_DIR = pathlib.Path("data/media/protocols_io_docs")

HEADERS = {"User-Agent": "LabiApp/1.0 (mailto:support@getlabi.app)"}
SIZE_CAP_BYTES = 20 * 1024 * 1024
TEXT_EXTRACTABLE = {"pdf", "docx"}


def extract_text_docx(path: pathlib.Path) -> str:
    doc = docx_lib.Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text_pdf(path: pathlib.Path) -> str:
    reader = pypdf.PdfReader(str(path))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    return "\n\n".join(parts)


def text_to_steps_raw(text: str) -> list:
    """Paragraph-split extracted document text into steps_raw dicts."""
    paras = re.split(r"\n{2,}", text)
    result = []
    for p in paras:
        p = re.sub(r"\s{2,}", " ", p.strip())
        if len(p) > 20:
            result.append({"title": "", "instruction": p})
    return result


def download(url: str, dest: pathlib.Path) -> int | None:
    try:
        with requests.get(url, headers=HEADERS, timeout=30, stream=True) as r:
            r.raise_for_status()
            dest.parent.mkdir(parents=True, exist_ok=True)
            total = 0
            with open(dest, "wb") as f:
                for chunk in r.iter_content(chunk_size=65536):
                    f.write(chunk)
                    total += len(chunk)
                    if total > SIZE_CAP_BYTES:
                        f.close()
                        dest.unlink()
                        return None
        return total
    except requests.RequestException as e:
        print(f"    [ERROR] {url}: {e}")
        if dest.exists():
            dest.unlink()
        return None


def sanitize(name: str) -> str:
    name = re.sub(r"[^\w.\-]", "_", name)
    return name[:150] or "file"


def process(dry_run: bool, limit: int | None = None) -> None:
    zero_step_files = []
    for norm_path in NORM_DIR.glob("*.json"):
        try:
            d = json.loads(norm_path.read_text())
        except Exception:
            continue
        if not d.get("steps") and "docs_checked" not in d:
            zero_step_files.append(norm_path)

    if limit:
        zero_step_files = zero_step_files[:limit]

    print(f"Zero-step protocols_io records to check: {len(zero_step_files)}")

    checked = no_cache = no_docs = downloaded = recovered = media_only = 0

    for norm_path in zero_step_files:
        checked += 1
        cache_path = CACHE_DIR / norm_path.name
        if not cache_path.exists():
            no_cache += 1
            continue

        try:
            cached = json.loads(cache_path.read_text())
        except Exception:
            no_cache += 1
            continue

        proto = cached.get("protocol") or cached
        docs = proto.get("documents") or []
        if not docs:
            no_docs += 1
            if not dry_run:
                d = json.loads(norm_path.read_text())
                d["docs_checked"] = True
                norm_path.write_text(json.dumps(d, ensure_ascii=False, indent=2))
            continue

        media_entries = []
        recovered_steps = None

        for doc in docs:
            url = doc.get("url", "")
            filename = sanitize(doc.get("ofn") or doc.get("filename") or f"doc_{doc.get('id')}")
            if not url:
                continue
            ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

            dest = MEDIA_DIR / norm_path.stem / filename
            if dry_run:
                print(f"  [WOULD FETCH] {norm_path.stem}: {filename} ({ext})")
                continue

            size = download(url, dest)
            if size is None:
                continue
            downloaded += 1

            media_entries.append({
                "type": "pdf" if ext == "pdf" else "document",
                "filename": filename,
                "local_path": str(dest),
                "source_url": url,
                "caption": "",
                "bytes": size,
            })

            if recovered_steps is None and ext in TEXT_EXTRACTABLE:
                try:
                    if ext == "docx" and DOCX_AVAILABLE:
                        text = extract_text_docx(dest)
                    elif ext == "pdf" and PDF_AVAILABLE:
                        text = extract_text_pdf(dest)
                    else:
                        text = ""
                    steps = text_to_steps_raw(text)
                    if len(steps) >= 2:
                        recovered_steps = steps
                except Exception as e:
                    print(f"    [EXTRACT ERROR] {filename}: {e}")

            time.sleep(0.3)

        if dry_run:
            continue

        d = json.loads(norm_path.read_text())
        d["docs_checked"] = True
        if media_entries:
            d["media"] = (d.get("media") or []) + media_entries
        if recovered_steps:
            canonical = []
            for i, s in enumerate(recovered_steps):
                canonical.append({
                    "step_id": i, "title": s["title"], "instruction": s["instruction"],
                    "substeps": [], "is_critical": False, "timers": [],
                })
            d["steps"] = canonical
            recovered += 1
            print(f"  [RECOVERED] {norm_path.stem}: {len(canonical)} steps from document")
        elif media_entries:
            media_only += 1

        norm_path.write_text(json.dumps(d, ensure_ascii=False, indent=2))

        # Mirror into raw for provenance.
        raw_path = RAW_DIR / norm_path.name
        if raw_path.exists():
            raw = json.loads(raw_path.read_text())
            if media_entries:
                raw["media"] = (raw.get("media") or []) + media_entries
            if recovered_steps:
                raw["steps_raw"] = recovered_steps
            raw_path.write_text(json.dumps(raw, ensure_ascii=False, indent=2))

    print(f"\nChecked: {checked} | No cache: {no_cache} | No documents: {no_docs}")
    print(f"Downloaded: {downloaded} | Recovered real steps: {recovered} | Media-only (no text recovered): {media_only}")
    if dry_run:
        print("(dry run — no files written)")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--limit", type=int)
    args = parser.parse_args()
    process(args.dry_run, args.limit)


if __name__ == "__main__":
    main()
